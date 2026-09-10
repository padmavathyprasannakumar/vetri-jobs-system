import os
import re
import json
import traceback

from groq import Groq


client = Groq(
    api_key=os.getenv("GROQ_API_KEY")
)


# =====================================================
# MODEL
# =====================================================
# "llama-3.1-8b-instant" was deprecated by Groq (June 2026).
# Calls to it now fail, which - because the old code swallowed
# every exception - showed up as a silent "Resume Score: 0"
# with no explanation. Groq's own migration guidance for that
# model is "openai/gpt-oss-20b", so that's what we use now.
# =====================================================

GROQ_MODEL = "openai/gpt-oss-20b"


# A couple of extra fallbacks in case GROQ_MODEL is ever
# retired/renamed again - each is tried in order until one
# succeeds, instead of the whole feature going dark.

GROQ_MODEL_FALLBACKS = [
    GROQ_MODEL,
    "openai/gpt-oss-120b",
    "llama-3.3-70b-versatile",
]


# =====================================================
# TEXT EXTRACTION
# =====================================================

def extract_resume_text(file_path):

    ext = os.path.splitext(file_path)[1].lower()


    if ext == ".docx":

        return _extract_text_from_docx(file_path)


    if ext == ".doc":

        # Legacy binary .doc isn't supported by python-docx or
        # PyMuPDF reliably. We still try PyMuPDF as a best effort
        # (it can sometimes read it), and fall back to an empty
        # string rather than crashing the whole upload.

        try:

            return _extract_text_from_pdf_like(file_path)

        except Exception:

            return ""


    # default: PDF (and anything PyMuPDF can open)

    return _extract_text_from_pdf_like(file_path)



def _extract_text_from_pdf_like(file_path):

    import fitz  # PyMuPDF - imported lazily so importing this

    # module doesn't hard-fail if PyMuPDF isn't installed yet.

    text = ""


    document = fitz.open(
        file_path
    )


    for page in document:

        text += page.get_text()


    document.close()


    return text



def _extract_text_from_docx(file_path):

    import docx  # python-docx


    document = docx.Document(
        file_path
    )


    parts = [
        p.text
        for p in document.paragraphs
        if p.text and p.text.strip()
    ]


    # Also pull text out of any tables (skills/experience are
    # sometimes laid out in table form in resume templates).

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text and cell.text.strip():

                    parts.append(cell.text)


    return "\n".join(parts)



# =====================================================
# JSON PARSING
# =====================================================
# Groq (like most chat models) will sometimes wrap JSON in
# ```json ... ``` fences, or add a short sentence before/after
# the object even when told not to. Instead of failing the
# whole analysis the moment that happens, we clean the text up
# and pull out the first {...} block before giving up.
# =====================================================

def _parse_ai_json(raw_text):

    text = (raw_text or "").strip()


    # strip ```json ... ``` or ``` ... ``` fences

    if text.startswith("```"):

        text = re.sub(r"^```[a-zA-Z]*\s*", "", text)

        text = re.sub(r"```\s*$", "", text)

        text = text.strip()


    try:

        return json.loads(text)

    except Exception:

        pass


    # fall back to grabbing the first {...} block in the text

    match = re.search(r"\{.*\}", text, re.DOTALL)


    if match:

        try:

            return json.loads(
                match.group(0)
            )

        except Exception:

            pass


    return None



# =====================================================
# AI RESUME ANALYSIS
# =====================================================

def analyse_resume_with_ai(file_path):


    resume_text = extract_resume_text(
        file_path
    )


    if not resume_text or not resume_text.strip():

        return {

            "resume_score": 0,

            "skills": [],

            "experience": [],

            "education": [],

            "certifications": [],

            "projects": [],

            "missing_information": [
                "Could not read any text from this file. "
                "Please upload a text-based PDF or DOCX "
                "(not a scanned image)."
            ],

            "job_categories": [],

            "text": resume_text or "",

        }


    # Groq/most LLM context windows comfortably fit a resume,
    # but trim extreme outliers so we never blow the token limit.

    resume_text = resume_text[:12000]


    prompt = f"""
You are an expert AI Resume Analyzer used inside a campus
placement platform.

Carefully read the resume text below and analyse it.

Return ONLY a single valid JSON object - no markdown, no code
fences, no commentary before or after it.

Resume:
---
{resume_text}
---

Required JSON format (use these exact keys):

{{
"resume_score": 0,
"skills": [],
"experience": [],
"education": [],
"certifications": [],
"projects": [],
"missing_information": [],
"job_categories": []
}}

Rules:
- "resume_score" is an integer 0-100 that reflects how complete
  and job-ready the resume is (structure, clarity, keywords,
  measurable outcomes, contact info, etc).
- "skills" is a flat list of technical/professional skills found.
- "experience" is a list of short strings, one per role
  (e.g. "Backend Intern at Acme Corp (2024-2025)").
- "education" is a list of short strings, one per degree/course.
- "certifications" and "projects" are lists of short strings.
- "missing_information" lists concrete, actionable gaps (e.g.
  "No professional summary", "No measurable project outcomes",
  "Missing contact email").
- "job_categories" suggests 2-5 suitable job titles/categories
  based on the resume content.
- If a section genuinely isn't present in the resume, return an
  empty list for it rather than guessing.
"""


    last_error = None


    for model_name in GROQ_MODEL_FALLBACKS:

        try:

            response = client.chat.completions.create(

                model=model_name,

                messages=[

                    {
                        "role": "user",
                        "content": prompt,
                    }

                ],

                temperature=0.2,

                response_format={
                    "type": "json_object"
                },

            )


            raw_content = response.choices[0].message.content


            parsed = _parse_ai_json(raw_content)


            if parsed:

                # make sure every expected key is present so the
                # frontend never has to guess

                return {

                    "resume_score": parsed.get("resume_score", 0),

                    "skills": parsed.get("skills", []),

                    "experience": parsed.get("experience", []),

                    "education": parsed.get("education", []),

                    "certifications": parsed.get("certifications", []),

                    "projects": parsed.get("projects", []),

                    "missing_information": parsed.get(
                        "missing_information", []
                    ),

                    "job_categories": parsed.get("job_categories", []),

                    "text": resume_text,

                }


            last_error = "AI returned a response that wasn't valid JSON"


        except Exception as e:

            # response_format=json_object isn't supported by every
            # model - retry that same model once without it before
            # moving on to the next fallback model.

            try:

                response = client.chat.completions.create(

                    model=model_name,

                    messages=[

                        {
                            "role": "user",
                            "content": prompt,
                        }

                    ],

                    temperature=0.2,

                )


                raw_content = response.choices[0].message.content

                parsed = _parse_ai_json(raw_content)


                if parsed:

                    return {

                        "resume_score": parsed.get("resume_score", 0),

                        "skills": parsed.get("skills", []),

                        "experience": parsed.get("experience", []),

                        "education": parsed.get("education", []),

                        "certifications": parsed.get(
                            "certifications", []
                        ),

                        "projects": parsed.get("projects", []),

                        "missing_information": parsed.get(
                            "missing_information", []
                        ),

                        "job_categories": parsed.get(
                            "job_categories", []
                        ),

                        "text": resume_text,

                    }


            except Exception as inner_e:

                last_error = str(inner_e)


            print(
                "AI Resume Analysis error on model",
                model_name,
                ":",
                e,
            )

            traceback.print_exc()


    # Every model attempt failed - surface *why* instead of a
    # silent zero score, so this is debuggable from the terminal
    # and from the Missing Information card in the UI.

    return {

        "resume_score": 0,

        "skills": [],

        "experience": [],

        "education": [],

        "certifications": [],

        "projects": [],

        "missing_information": [
            f"AI analysis failed: {last_error or 'unknown error'}"
        ],

        "job_categories": [],

        "text": resume_text,

    }


# =====================================================
# RESUME ANALYSIS WRAPPER
# (kept for backwards compatibility with existing callers
# that import "analyze_resume" instead of
# "analyse_resume_with_ai")
# =====================================================

def analyze_resume(file_path):

    return analyse_resume_with_ai(
        file_path
    )


# =====================================================
# ATS-FRIENDLINESS CHECK
# Analyses a resume specifically for Applicant Tracking
# System (ATS) compatibility - formatting, structure,
# keyword usage - separate from the general quality/score
# analysis above, and returns concrete rewrite suggestions.
# =====================================================

def analyze_ats_friendliness(file_path, target_role=None):

    resume_text = extract_resume_text(
        file_path
    )

    if not resume_text or not resume_text.strip():

        return {
            "ats_score": 0,
            "issues": [
                "Could not read any text from this file - if it's a "
                "scanned image or has heavy graphics/columns, that "
                "itself is a major ATS problem. Use a simple, "
                "single-column text-based PDF or DOCX."
            ],
            "suggestions": [],
            "rewritten_bullets": [],
        }

    resume_text = resume_text[:12000]

    role_line = (
        f'Target role: "{target_role}".'
        if target_role else
        "No specific target role given - assess generally."
    )

    prompt = f"""
You are an ATS (Applicant Tracking System) compatibility expert.

Analyse the resume text below purely for how well it would parse and
rank in an automated ATS, NOT general resume quality.

{role_line}

Resume:
---
{resume_text}
---

Return ONLY a single valid JSON object - no markdown, no commentary:

{{
"ats_score": 0,
"issues": [],
"suggestions": [],
"rewritten_bullets": []
}}

Rules:
- "ats_score" is an integer 0-100 for ATS-parseability (structure,
  standard section headings, no tables/columns/graphics/images,
  standard fonts, no headers/footers with contact info, keyword
  presence, simple bullet formatting, file structure).
- "issues" lists specific, concrete ATS problems this resume
  actually appears to have (e.g. "Contact info may be in a header,
  which some ATS parsers skip", "Uses a two-column layout which can
  scramble reading order", "Missing a standard 'Skills' section
  heading"). Empty list if none found.
- "suggestions" lists concrete fixes tied to the issues above (e.g.
  "Move phone/email into the main body, not a header/footer",
  "Switch to a single-column, reverse-chronological layout",
  "Add a dedicated 'Skills' section listing exact keywords from job
  postings you're targeting").
- "rewritten_bullets" contains 2-4 example rewrites: take a weak
  bullet point from the resume text (marked "before") and rewrite it
  ATS-friendly (marked "after") - active verb, quantified outcome,
  relevant keywords, no special characters/icons. Format each as a
  string "Before: ... | After: ...". Empty list if the resume has no
  bullet points to rewrite.
"""

    last_error = None

    for model_name in GROQ_MODEL_FALLBACKS:

        try:

            response = client.chat.completions.create(

                model=model_name,

                messages=[
                    {"role": "user", "content": prompt}
                ],

                temperature=0.3,

                response_format={"type": "json_object"},

            )

            parsed = _parse_ai_json(
                response.choices[0].message.content
            )

            if parsed:

                return {
                    "ats_score": parsed.get("ats_score", 0),
                    "issues": parsed.get("issues", []),
                    "suggestions": parsed.get("suggestions", []),
                    "rewritten_bullets": parsed.get(
                        "rewritten_bullets", []
                    ),
                }

            last_error = "AI returned invalid JSON"

        except Exception as e:

            last_error = str(e)

            print("ATS check error on model", model_name, ":", e)

            continue

    return {
        "ats_score": 0,
        "issues": [
            f"ATS analysis failed: {last_error or 'unknown error'}"
        ],
        "suggestions": [],
        "rewritten_bullets": [],
    }
